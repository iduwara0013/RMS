from __future__ import annotations

import io
import re
from collections import defaultdict
from typing import Any

import pymupdf
from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from PIL import Image
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P

try:
    import pytesseract
except ImportError:  # OCR stays optional at runtime.
    pytesseract = None


MAX_FILE_SIZE = 5 * 1024 * 1024
PARSER_VERSION = "layout-v1"
ALLOWED_EXTENSIONS = {"pdf", "docx"}

SECTION_ALIASES: dict[str, set[str]] = {
    "summary": {
        "profile", "professional profile", "personal profile", "summary",
        "professional summary", "career summary", "objective", "career objective",
        "about me", "about",
    },
    "skills": {
        "skills", "technical skills", "key skills", "core skills", "competencies",
        "core competencies", "technical competencies", "expertise", "technologies",
    },
    "experience": {
        "experience", "work experience", "professional experience", "employment history",
        "career history", "work history", "employment", "professional background",
    },
    "education": {
        "education", "academic background", "academic qualifications", "educational qualifications",
        "qualifications", "education and qualifications", "academic history",
    },
    "projects": {
        "projects", "project experience", "selected projects", "academic projects",
        "personal projects", "key projects", "portfolio",
    },
    "certifications": {
        "certification", "certifications", "certificate", "certificates", "professional certifications", "licenses",
        "licenses and certifications", "courses", "training",
    },
    "languages": {
        "languages", "language", "language proficiency", "language skills",
    },
    "publications": {
        "publications", "research", "publications and research", "research and publications",
    },
}

ALL_HEADINGS = {alias: section for section, aliases in SECTION_ALIASES.items() for alias in aliases}


app = FastAPI(title="CPSTL CV Parser", version=PARSER_VERSION, docs_url=None, redoc_url=None)


def clean(value: str) -> str:
    value = value.replace("\u00a0", " ").replace("\r", "")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def normalized_heading(value: str) -> str:
    value = clean(value).lower().replace("&", "and")
    value = re.sub(r"[^a-z ]", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def heading_for(value: str) -> str | None:
    candidate = normalized_heading(value)
    if candidate in ALL_HEADINGS:
        return ALL_HEADINGS[candidate]
    for alias, section in ALL_HEADINGS.items():
        if candidate.startswith(alias + " ") and len(candidate) <= len(alias) + 18:
            return section
    return None


def pdf_blocks(content: bytes) -> tuple[list[dict[str, Any]], bool, bool]:
    blocks: list[dict[str, Any]] = []
    used_ocr = False
    document = pymupdf.open(stream=content, filetype="pdf")

    for page_number, page in enumerate(document):
        page_blocks: list[dict[str, Any]] = []
        data = page.get_text("dict", sort=False)
        for block in data.get("blocks", []):
            if block.get("type") != 0:
                continue
            lines: list[str] = []
            sizes: list[float] = []
            bold = False
            for line in block.get("lines", []):
                text = "".join(span.get("text", "") for span in line.get("spans", []))
                if clean(text):
                    lines.append(clean(text))
                for span in line.get("spans", []):
                    sizes.append(float(span.get("size", 0)))
                    bold = bold or "bold" in str(span.get("font", "")).lower()
            text = clean("\n".join(lines))
            if text:
                x0, y0, x1, y1 = block.get("bbox", (0, 0, 0, 0))
                page_blocks.append({
                    "text": text, "page": page_number, "x0": float(x0), "y0": float(y0),
                    "x1": float(x1), "y1": float(y1), "page_width": float(page.rect.width),
                    "font_size": max(sizes, default=0), "bold": bold,
                })

        readable = sum(len(block["text"]) for block in page_blocks)
        if readable < 40 and pytesseract is not None:
            try:
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                ocr_text = clean(pytesseract.image_to_string(image))
                if ocr_text:
                    page_blocks.append({
                        "text": ocr_text, "page": page_number, "x0": 0.0, "y0": 0.0,
                        "x1": float(page.rect.width), "y1": float(page.rect.height),
                        "page_width": float(page.rect.width), "font_size": 0.0, "bold": False,
                    })
                    used_ocr = True
            except Exception:
                pass

        blocks.extend(page_blocks)

    document.close()
    blocks.sort(key=lambda block: (block["page"], round(block["x0"] / max(block["page_width"], 1), 1), block["y0"]))
    return blocks, used_ocr, bool(blocks)


def docx_blocks(content: bytes) -> tuple[list[dict[str, Any]], bool, bool]:
    document = Document(io.BytesIO(content))
    blocks: list[dict[str, Any]] = []
    position = 0
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)
            text = clean(paragraph.text)
            if text:
                style = (paragraph.style.name or "").lower() if paragraph.style else ""
                bold = any(run.bold for run in paragraph.runs if run.text.strip())
                blocks.append({
                    "text": text, "page": 0, "x0": 0.0, "y0": float(position), "x1": 1.0,
                    "y1": float(position + 1), "page_width": 1.0,
                    "font_size": 16.0 if "heading" in style or "title" in style else 11.0,
                    "bold": bold or "heading" in style or "title" in style,
                })
            position += 1
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)
            for row in table.rows:
                column_count = max(len(row.cells), 1)
                seen_cells: set[int] = set()
                for column, cell in enumerate(row.cells):
                    cell_key = id(cell._tc)
                    if cell_key in seen_cells:
                        continue
                    seen_cells.add(cell_key)
                    text = clean("\n".join(paragraph.text for paragraph in cell.paragraphs))
                    if not text:
                        continue
                    cell_paragraphs = [paragraph for paragraph in cell.paragraphs if clean(paragraph.text)]
                    bold = any(run.bold for paragraph in cell_paragraphs for run in paragraph.runs if run.text.strip())
                    has_heading_style = any(
                        paragraph.style and ("heading" in paragraph.style.name.lower() or "title" in paragraph.style.name.lower())
                        for paragraph in cell_paragraphs
                    )
                    blocks.append({
                        "text": text, "page": 0, "x0": column / column_count,
                        "y0": float(position), "x1": (column + 1) / column_count,
                        "y1": float(position + 1), "page_width": 1.0,
                        "font_size": 16.0 if has_heading_style else 11.0,
                        "bold": bold or has_heading_style,
                    })
                position += 1
    return blocks, False, bool(blocks)


def group_by_layout(blocks: list[dict[str, Any]]) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = defaultdict(list)
    headings: list[dict[str, Any]] = []
    body_sizes = [block["font_size"] for block in blocks if block["font_size"] > 0]
    median_size = sorted(body_sizes)[len(body_sizes) // 2] if body_sizes else 11.0

    for block in blocks:
        first_line = block["text"].split("\n", 1)[0]
        section = heading_for(first_line)
        looks_like_heading = bool(section) and (
            block["bold"] or block["font_size"] >= median_size or first_line.isupper()
        )
        if looks_like_heading:
            heading = dict(block)
            heading["section"] = section
            headings.append(heading)
            remainder = clean(block["text"][len(first_line):])
            if remainder:
                sections[section].append(remainder)

    for block in blocks:
        first_line = block["text"].split("\n", 1)[0]
        if heading_for(first_line) and any(
            heading["page"] == block["page"] and heading["x0"] == block["x0"] and heading["y0"] == block["y0"]
            for heading in headings
        ):
            continue

        candidates = []
        block_center = (block["x0"] + block["x1"]) / 2
        for heading in headings:
            if heading["page"] != block["page"] or heading["y0"] > block["y0"] + 2:
                continue
            heading_center = (heading["x0"] + heading["x1"]) / 2
            horizontal_distance = abs(block_center - heading_center)
            left_edges_align = abs(block["x0"] - heading["x0"]) <= block["page_width"] * 0.12
            centers_align = horizontal_distance <= block["page_width"] * 0.28
            heading_is_wide = (heading["x1"] - heading["x0"]) >= block["page_width"] * 0.65
            if left_edges_align or centers_align or heading_is_wide:
                candidates.append((block["y0"] - heading["y0"], horizontal_distance, heading))
        if candidates:
            _, _, heading = min(candidates, key=lambda item: (item[0], item[1]))
            sections[heading["section"]].append(block["text"])

    return sections


def lines(values: list[str]) -> list[str]:
    result: list[str] = []
    for value in values:
        result.extend(clean(line) for line in value.splitlines() if clean(line))
    return result


def unique(values: list[str], limit: int) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        value = clean(value).strip("-–—•·;,. ")
        key = value.casefold()
        if len(value) >= 2 and key not in seen:
            seen.add(key)
            result.append(value)
    return result[:limit]


def skill_items(values: list[str]) -> list[str]:
    items: list[str] = []
    category = re.compile(r"^(?:front(?:end|ed)(?:\s*\/\s*mobile)?|backend|mobile|cloud|devops|databases?(?:\s*&\s*tools)?|tools|frameworks?|programming languages?|ai(?:\s*\/\s*emerging technologies)?)\s*[-:–—]", re.I)
    source: list[str] = []
    for line in lines(values):
        if source and category.match(source[-1]) and not category.match(line):
            source[-1] += " " + line
        else:
            source.append(line)
    for line in source:
        line = re.sub(r"^(?:front(?:end|ed)(?:\s*\/\s*mobile)?|backend|mobile|cloud|devops|databases?(?:\s*&\s*tools)?|tools|frameworks?|programming languages?|ai(?:\s*\/\s*emerging technologies)?)\s*[-:–—]\s*", "", line, flags=re.I)
        items.extend(re.split(r",(?![^()]*\))|[;•|]", line))
    return unique(items, 30)


def project_items(values: list[str]) -> list[str]:
    source = lines(values)
    projects: list[str] = []
    current: list[str] = []
    for line in source:
        starts_project = bool(re.search(r"\[(?:git(?:hub)?|project)\s*link\]", line, re.I))
        if starts_project and current:
            projects.append(" ".join(current))
            current = []
        current.append(line)
    if current:
        projects.append(" ".join(current))
    return unique(projects, 12)


def fallback_fields(blocks: list[dict[str, Any]], result: dict[str, Any]) -> None:
    all_lines = lines([block["text"] for block in blocks])
    if not result["education"]:
        for index, line in enumerate(all_lines):
            if re.search(r"\b(B\.?SC|M\.?SC|BACHELOR|MASTER|DIPLOMA|PHD|DEGREE|HND)\b", line, re.I):
                result["education"].append(" · ".join(all_lines[index:index + 3]))
    if not result["certifications"]:
        result["certifications"] = unique([
            line for line in all_lines
            if re.search(r"\b(CCNA|COMPTIA|CISCO|ORACLE|AWS|AZURE|CERTIFICATION|CERTIFIED)\b", line, re.I)
        ], 12)
    if not result["skills"]:
        category_lines = [
            line for line in all_lines
            if re.match(r"^(frontend|fronted|backend|mobile|cloud|devops|databases?|tools|frameworks?|programming|ai)\b", line, re.I)
        ]
        result["skills"] = skill_items(category_lines)


def normalize_profile(blocks: list[dict[str, Any]], used_ocr: bool, has_text: bool) -> dict[str, Any]:
    sections = group_by_layout(blocks)
    result: dict[str, Any] = {
        "summary": clean("\n".join(sections.get("summary", [])))[:2000] or None,
        "skills": skill_items(sections.get("skills", [])),
        "experience": unique(lines(sections.get("experience", [])), 20),
        "education": unique(lines(sections.get("education", [])), 16),
        "projects": project_items(sections.get("projects", [])),
        "certifications": unique(lines(sections.get("certifications", [])), 16),
        "languages": unique(re.split(r"[,;•|]", " ".join(lines(sections.get("languages", [])))), 12),
    }
    fallback_fields(blocks, result)

    populated = sum(bool(result[key]) for key in ("summary", "skills", "experience", "education", "projects", "certifications", "languages"))
    recognized = sum(bool(sections.get(key)) for key in SECTION_ALIASES)
    confidence = min(96, 35 + populated * 7 + recognized * 3 + (0 if used_ocr else 7)) if has_text else 0
    review_status = "Pending HR review" if confidence >= 65 else "Needs review"
    message = (
        "Layout-aware extraction completed. Confirm important details against the original CV."
        if confidence >= 65 else
        "Only part of this CV could be classified automatically. Review the original document."
    )

    result.update({
        "parse_status": "Parsed" if has_text else "Needs review",
        "parse_message": message,
        "confidence_score": confidence,
        "review_status": review_status,
        "parser_version": PARSER_VERSION,
        "parser_metadata": {
            "layout_blocks": len(blocks),
            "recognized_sections": recognized,
            "used_ocr": used_ocr,
        },
    })
    return result


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok", "parser_version": PARSER_VERSION}


@app.post("/extract")
async def extract(file: UploadFile = File(...)) -> dict[str, Any]:
    filename = file.filename or "cv"
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=422, detail="Only PDF and DOCX CVs can be processed.")

    content = await file.read(MAX_FILE_SIZE + 1)
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="CV exceeds the 5 MB processing limit.")

    try:
        if extension == "pdf":
            blocks, used_ocr, has_text = pdf_blocks(content)
        else:
            blocks, used_ocr, has_text = docx_blocks(content)
        return normalize_profile(blocks, used_ocr, has_text)
    except Exception as exception:
        raise HTTPException(status_code=422, detail="The CV could not be parsed safely.") from exception
